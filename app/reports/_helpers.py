"""Shared helpers for report rendering.

Conventions:
    columns = [{"field": str, "label": str, "fmt"?: str, "width"?: int, "align"?: str}]
    rows    = [dict]
    summary = {key: value} — flat, value rendered as-is
    extra_sheets = [{"title": str, "columns": columns, "rows": rows}]
"""

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


NUMBER_FORMATS = {
    "int":      "#,##0",
    "money":    "#,##0.00",
    "money1":   "#,##0.0",
    "pct":      "0.00%",
    "date":     "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm",
}

_HEADER_FILL = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
_HEADER_FONT = Font(bold=True)


def _coerce_value(value, fmt):
    if value is None or value == "":
        return None
    if fmt in ("date", "datetime") and isinstance(value, str):
        try:
            if "T" in value or " " in value:
                return datetime.fromisoformat(value.replace("Z", ""))
            return date.fromisoformat(value)
        except ValueError:
            return value
    return value


def _write_sheet(ws, columns, rows):
    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col.get("label") or col["field"])
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        width = col.get("width")
        if width:
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col in enumerate(columns, start=1):
            field = col["field"]
            fmt = col.get("fmt")
            value = _coerce_value(row.get(field), fmt)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if fmt and fmt in NUMBER_FORMATS:
                cell.number_format = NUMBER_FORMATS[fmt]
            align = col.get("align")
            if align:
                cell.alignment = Alignment(horizontal=align)
            elif fmt in ("int", "money", "money1", "pct"):
                cell.alignment = Alignment(horizontal="right")

    ws.freeze_panes = "A2"


def _write_summary_sheet(ws, summary, summary_labels, title="สรุป"):
    ws.cell(row=1, column=1, value=title).font = _HEADER_FONT
    row = 2
    for key, value in summary.items():
        label = summary_labels.get(key, key) if summary_labels else key
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 24


def render_xlsx_report(
    sheet_title,
    columns,
    rows,
    summary=None,
    summary_labels=None,
    summary_title="สรุป",
    extra_sheets=None,
):
    """Build an XLSX workbook in memory; return bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Data")[:31]
    _write_sheet(ws, columns, rows)

    if summary:
        s_ws = wb.create_sheet(title=summary_title[:31])
        _write_summary_sheet(s_ws, summary, summary_labels, title=summary_title)

    for extra in extra_sheets or []:
        e_ws = wb.create_sheet(title=(extra.get("title") or "Sheet")[:31])
        _write_sheet(e_ws, extra["columns"], extra.get("rows") or [])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_json_default(data):
    """Default JSON shape: {meta: summary, rows: rows}."""
    return {"meta": data.get("summary") or {}, "rows": data.get("rows") or []}
